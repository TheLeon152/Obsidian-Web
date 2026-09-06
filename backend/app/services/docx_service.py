from pathlib import Path

from docx import Document


class DocxService:

    def render_html(self, path: Path) -> str:
        document = Document(path)

        html_parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if not text:
                continue

            style_name = paragraph.style.name.lower()

            if "heading 1" in style_name:
                html_parts.append(
                    f"<h1>{self._escape(text)}</h1>"
                )
            elif "heading 2" in style_name:
                html_parts.append(
                    f"<h2>{self._escape(text)}</h2>"
                )
            elif "heading 3" in style_name:
                html_parts.append(
                    f"<h3>{self._escape(text)}</h3>"
                )
            else:
                html_parts.append(
                    f"<p>{self._escape(text)}</p>"
                )

        for table in document.tables:
            html_parts.append(self._render_table(table))

        return "\n".join(html_parts)

    @staticmethod
    def _escape(text: str) -> str:
        return (
            text
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
            .replace("'", "&#39;")
        )

    def _render_table(self, table) -> str:
        rows: list[str] = []

        for row in table.rows:
            cells = [
                f"<td>{self._escape(cell.text)}</td>"
                for cell in row.cells
            ]

            rows.append(
                f"<tr>{''.join(cells)}</tr>"
            )

        return (
            '<table class="docx-viewer-table">'
            f"<tbody>{''.join(rows)}</tbody>"
            "</table>"
        )